import os
import json
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REP = os.path.join(ROOT, "reports")
FIG = os.path.join(ROOT, "slides")

def _box(slide, lines, left, top, width, height, size, bold=False, color=RGBColor(0x22, 0x22, 0x22)):
    tf = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height)).text_frame
    tf.word_wrap = True
    for n, line in enumerate(lines):
        p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
        p.text = line
        p.font.size, p.font.bold, p.font.color.rgb = Pt(size), bold, color
        p.space_after = Pt(size * 0.5)

def title(slide, text):
    _box(slide, [text], 0.5, 0.3, 12.3, 0.9, 30, True, RGBColor(0x1F, 0x3A, 0x5F))

def bullets(slide, lines, left, top, width, height, size=16):
    _box(slide, ["• " + l for l in lines], left, top, width, height, size)

def table(slide, rows, left, top, width, height, size):
    t = slide.shapes.add_table(len(rows), len(rows[0]), Inches(left), Inches(top), Inches(width), Inches(height)).table
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
            t.cell(i, j).text_frame.paragraphs[0].font.size = Pt(size)

def pic(slide, name, left, top, width):
    slide.shapes.add_picture(os.path.join(FIG, name), Inches(left), Inches(top), width=Inches(width))

NAMES = {"APP_EXT_MEAN": "Average external score", "APP_ORGANIZATION_TYPE": "Employer type", "APP_EXT_MIN": "Lowest external score",
         "APP_EXT_MAX": "Highest external score", "APP_EXT_SOURCE_3": "External score 3", "APP_EXT_SOURCE_2": "External score 2",
         "APP_CREDIT_TO_ANNUITY": "Loan term (credit / annuity)", "BUR_DEBT_RATIO_MAX": "Highest debt share on a bureau loan",
         "APP_OCCUPATION_TYPE": "Occupation", "INS_RECENT_LATE_RATE": "Late instalments, last 12 months",
         "PREV_CREDIT_GOODS_RATIO_MEAN": "Past loans, credit / goods price", "APP_CREDIT_TO_GOODS": "Credit / goods price"}

def importance_chart():
    imp = pd.read_csv(os.path.join(REP, "lgbm_importance.csv")).head(len(NAMES)).iloc[::-1]
    fig, ax = plt.subplots(figsize=(7, 6.6))
    ax.barh([NAMES.get(f, f) for f in imp.feature], imp.gain / imp.gain.max(), color="#3b6ea5")
    ax.set_title("LightGBM, top 12 inputs (importance, top = 1)", fontsize=13)
    ax.tick_params(labelsize=12)
    ax.spines[["top", "right"]].set_visible(False)
    plt.savefig(os.path.join(FIG, "lgbm_importance_named.png"), bbox_inches="tight", dpi=150)
    plt.close(fig)

def footer(slide, n):
    _box(slide, [f"Home Credit Indonesia · Data Scientist assessment · Nauval Zulfikar · {n}/5"], 0.5, 7.0, 12.3, 0.4, 10, color=RGBColor(0x80, 0x80, 0x80))

def main():
    m1 = json.load(open(os.path.join(REP, "lgbm_metrics.json")))
    m2 = json.load(open(os.path.join(REP, "logreg_metrics.json")))
    m3 = json.load(open(os.path.join(REP, "lgbm_noext_metrics.json")))
    iv = pd.read_csv(os.path.join(REP, "logreg_iv.csv"))
    sc = pd.read_csv(os.path.join(REP, "scorecard.csv"))
    ev = json.load(open(os.path.join(REP, "evaluate_summary.json")))
    vs = json.load(open(os.path.join(REP, "value_summary.json")))
    rc = pd.read_csv(os.path.join(REP, "reject_curve.csv")); rc = rc[rc.model == "LightGBM"]
    val = pd.read_csv(os.path.join(REP, "value.csv")); val = val[val.model == "LightGBM"]
    cut = rc.merge(val[["reject_pct", "break_even_ratio"]], on="reject_pct")
    cut = cut[cut.reject_pct.isin([0.05, 0.10, 0.15, 0.20, 0.30])]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "Predicting credit default, data and risk signals")
    bullets(slide, [
        f"{m1['n_train']:,} past loans, {m1['default_rate']:.1%} defaulted; {m1['n_test']:,} new applicants to score",
        "7 linked tables (application, credit bureau (+monthly balance), previous Home Credit loans, instalments, POS & credit-card balances)",
        "Strongest signals are the external credit scores; for EXT_SOURCE_2 the default rate falls from 18.4% in the lowest decile to 3.0% in the highest",
        "Younger applicants default more, 12.3% at 20-25 vs 4.9% at 60-70",
        "Clients with the 'pensioner' placeholder (DAYS_EMPLOYED = 365243) default at 5.4% vs 8.7% for the employed, kept as a flag",
        "Imbalanced target → judged by AUC / Gini / KS, not accuracy"
    ], 0.5, 1.2, 12.3, 3.0, 15)
    pic(slide, "eda.png", 1.4, 4.2, 10.5)
    footer(slide, 1)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "Workflow & business-logic features")
    bullets(slide, [
        "Raw CSVs → 4 feature scripts → 2 models → validation → scoring (one command, run_all.sh)",
        "Affordability, credit/income, annuity/income, loan term (credit/annuity)",
        "Stability, age, years employed, employment-to-age ratio",
        "Bureau, active loans, total debt ratio, worst days-past-due",
        "Past Home Credit behaviour, refusal rate, recent refusals, late-payment rate, under-payment, card utilisation",
        f"{m1['n_features']} features for LightGBM; {m2['n_selected']} for the scorecard (IV filter + correct risk direction only)",
        f"Without the external scores LightGBM still reaches AUC {m3['oof_auc']:.3f}, above the best simple rule ({ev['best_baseline_auc']:.3f}), so the built features carry their own signal",
        "Gender excluded from both models (fair lending)"
    ], 0.5, 1.3, 6.3, 5.6, 15)
    importance_chart()
    pic(slide, "lgbm_importance_named.png", 6.9, 1.2, 6.0)
    footer(slide, 2)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "Two models, accuracy vs. explainability")
    pic(slide, "compare.png", 0.4, 1.2, 7.6)
    table(slide, [
        ["Metric", "LightGBM", "Scorecard (LR)"],
        ["OOF AUC", f"{m1['oof_auc']:.3f}", f"{m2['oof_auc']:.3f}"],
        ["Gini", f"{m1['oof_gini']:.3f}", f"{m2['oof_gini']:.3f}"],
        ["KS", f"{m1['oof_ks']:.3f}", f"{m2['oof_ks']:.3f}"],
        ["Fold AUC std", f"{m1['auc_std']:.4f}", f"{m2['auc_std']:.4f}"],
        ["AUC on own training data", f"{m1['train_auc_last_fold']:.3f}", f"{m2['train_auc_full']:.3f}"],
        ["PSI train→test", f"{m1['psi_train_vs_test']:.4f}", f"{m2['psi_train_vs_test']:.4f}"],
    ], 8.3, 1.4, 4.6, 3.2, 13)
    bullets(slide, [
        f"Both beat the best simple rule (average of external scores, AUC {ev['best_baseline_auc']:.3f}); both stable across folds (std < {max(m1['auc_std'], m2['auc_std']) + 0.0005:.3f}); scores stable train→test (PSI < 0.01) but the loan-term input shifted",
        f"LightGBM +{m1['oof_auc']-m2['oof_auc']:.3f} AUC, but train {m1['train_auc_last_fold']:.2f} vs valid {m1['oof_auc']:.2f} → watch overfit",
        f"Scorecard has no overfit, every coefficient has the expected sign, points {m2['score_min']:.0f} to {m2['score_max']:.0f}"
    ], 0.5, 4.6, 7.5, 2.3, 15)
    footer(slide, 3)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "Scorecard, every decision comes with a reason")
    meaning = {"APP_EXT_MEAN": "average external credit score", "APP_EXT_MIN": "worst external credit score",
               "APP_EXT_SOURCE_3": "external score 3", "APP_EXT_SOURCE_2": "external score 2", "APP_EXT_SOURCE_1": "external score 1",
               "BUR_DAYS_CREDIT_MEAN": "how recent bureau loans are", "APP_EMPLOYED_YEARS": "years in current job",
               "BUR_TOTAL_DEBT_RATIO": "outstanding / total bureau debt", "APP_EMPLOYED_TO_AGE": "share of life employed",
               "BUR_ACTIVE_MEAN": "share of bureau loans still active"}
    top = iv[iv.selected].sort_values("iv", ascending=False).head(8)
    table(slide, [["Top variables (by IV)", "Meaning", "IV"]] +
          [[r.name, meaning.get(r.name, ""), f"{r.iv:.2f}"] for r in top.itertuples()], 0.5, 1.3, 6.4, 3.4, 11)
    late = sc[(sc.Variable == "INS_RECENT_LATE_RATE") & (sc.Bin != "Special")]
    miss_r = next(late[late.Bin == "Missing"].itertuples())
    def band(b):
        if b == "Missing":
            return "none in last 12 months"
        lo, hi = (float(x) for x in b.strip("[()").split(", "))
        return f"< {hi:.0%}" if lo == float("-inf") else f"≥ {lo:.0%}" if hi == float("inf") else f"{lo:.0%}-{hi:.0%}"
    _box(slide, ["Example, INS_RECENT_LATE_RATE (late instalments, last 12 months)"], 7.2, 1.2, 5.7, 0.4, 12, True)
    table(slide, [["Late share", "Clients", "Default rate", "Points"]] +
          [[band(r.Bin), f"{r._5:.0%}", f"{r._8:.1%}", f"{r.Points:.1f}"] for r in late.itertuples()], 7.2, 1.9, 5.7, 2.8, 12)
    bullets(slide, [
        f"{m2['n_selected']} of {m2['n_candidates']} candidates kept, with information value ≥ 0.02, then only variables whose effect has the expected sign ({', '.join(m2['dropped_wrong_sign']) or 'none'} dropped)",
        "Score = sum of points per variable; 600 points = 50:1 good:bad odds, every 20 points doubles the odds",
        "Reason codes for a rejection = the variables where the applicant lost the most points vs. the best bin",
        f"Missing values form their own bin, scored by their own default rate (no instalments in the last 12 months, {miss_r._5:.0%} of clients, {miss_r._8:.1%} default, {miss_r.Points:.1f} points)",
    ], 0.5, 5.3, 12.3, 1.7, 14)
    footer(slide, 4)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "Business use & next steps")
    table(slide, [["Reject riskiest", "Default rate of approved", "Defaults avoided", "Break-even ratio"]] +
          [[f"{r.reject_pct:.0%}", f"{r.approved_bad_rate:.1%}", f"{r.caught_pct:.0%}", f"{r.break_even_ratio:.2f}"] for r in cut.itertuples()],
          0.5, 1.4, 6.0, 2.6, 14)
    bullets(slide, [
        f"Today {m1['default_rate']:.1%} of loans default",
        f"Refusing the riskiest 10% pays off only if a default loses more than {vs['lgbm_at_10']['break_even_ratio']:.1f}x what a good loan earns, per unit of credit",
        "Break-even ratio = credit of good clients refused ÷ credit of defaulters refused; the bank plugs in its own margin",
        "Past refusals have no outcome → confirm with a small live test"
    ], 0.5, 4.3, 6.0, 2.6, 14)
    bullets(slide, [
        "LightGBM for ranking; the scorecard where reason codes are needed (adverse-action letters, OJK audit); recalibrate on recent loans before pricing",
        "Grey zone, auto-approve low risk, auto-reject top risk, manual review in the middle",
        "Monthly, PSI of the score and top-10 inputs (watch > 0.1, act > 0.25) and early default; quarterly, Gini, calibration and refusal rate by gender and age",
        f"Limits, no application dates so no time-based test; test set is {ev['test_cash_share']:.0%} cash loans (cash-loan AUC {ev['lgbm_auc_cash']:.3f}); heavy reliance on external scores",
        f"Loan terms are shorter in test (PSI {vs['feature_psi_max']:.1f}) though the score stays stable; under-30s are refused far more than their risk explains",
        "Next, shadow run, go live once the shift is explained, agree the cut-off with the business, validate on dated data"
    ], 6.9, 1.4, 6.0, 5.4, 14)
    footer(slide, 5)

    prs.save(os.path.join(ROOT, "HomeCredit_Assessment_Nauval.pptx"))
    print("saved", len(prs.slides), "slides")

if __name__ == "__main__":
    main()
